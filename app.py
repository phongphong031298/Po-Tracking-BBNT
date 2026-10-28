"""
app.py
PO Tracking App - Tra cứu, Nhập liệu & Sửa trực tiếp
"""
import streamlit as st
import pandas as pd
from datetime import datetime, date
from pathlib import Path
import data_manager as dm
import openpyxl
import os
import re

# Cấu hình trang
st.set_page_config(
    page_title="PO Tracking - Quản lý đơn hàng",
    page_icon="📦",
    layout="wide",
)

# =========================================================
# CẤU HÌNH FILE MẶC ĐỊNH
# =========================================================
DEFAULT_FILE = "data/theo_doi_po.xlsx"

if "file_path" not in st.session_state:
    st.session_state.file_path = DEFAULT_FILE

# =========================================================
# HÀM LÀM SẠCH KEY (QUAN TRỌNG ĐỂ SỬA LỖI)
# =========================================================
def clean_key(text):
    """Chuyển chuỗi thành dạng an toàn để làm key cho Streamlit (chỉ giữ chữ và số)"""
    return re.sub(r'[^a-zA-Z0-9]', '', str(text))

# =========================================================
# SIDEBAR
# =========================================================
st.sidebar.header("⚙️ Cấu hình dữ liệu")
uploaded = st.sidebar.file_uploader("Tải file Excel khác (nếu muốn)", type=["xlsx"])

if uploaded:
    save_path = Path("data") / uploaded.name
    save_path.parent.mkdir(exist_ok=True)
    with open(save_path, "wb") as f:
        f.write(uploaded.getbuffer())
    st.session_state.file_path = str(save_path)
    st.sidebar.success(f"Đã tải: {uploaded.name}")
    st.session_state.pop("all_sheets", None)

file_path = st.session_state.file_path

if not Path(file_path).exists():
    st.error(f"Không tìm thấy file dữ liệu tại: {file_path}")
    st.info("Vui lòng upload file Excel ở thanh bên trái để bắt đầu.")
    st.stop()

st.sidebar.info(f"📁 File đang dùng: **{Path(file_path).name}**")

if "all_sheets" not in st.session_state:
    with st.spinner("Đang tải dữ liệu..."):
        st.session_state.all_sheets = dm.load_all_sheets(file_path)
        st.session_state.visible_sheets = dm.get_visible_sheets(file_path)

all_data = st.session_state.all_sheets
visible_sheets = st.session_state.visible_sheets

# =========================================================
# TABS
# =========================================================
tab_tra_cuu, tab_loc_ngay, tab_tong_hop = st.tabs([
    "🔍 Tra cứu & Nhập liệu", 
    "📅 Lịch sử giao hàng theo Ngày", 
    "📊 Tổng hợp KPI"
])

# ---------------------------------------------------------
# TAB 1: TRA CỨU & NHẬP LIỆU
# ---------------------------------------------------------
with tab_tra_cuu:
    st.header("🔍 Tra cứu & Nhập liệu theo Công ty")
    
    col1, col2 = st.columns([1, 2])
    with col1:
        cong_ty = st.selectbox("Chọn Công ty (Sheet)", visible_sheets)
    with col2:
        search_kw = st.text_input("Lọc nhanh Mã SP / Tên SP", placeholder="Gõ để tìm...")

    df_company = all_data.get(cong_ty)
    
    if isinstance(df_company, pd.DataFrame) and not df_company.empty and "Mã SP" in df_company.columns:
        df_si = df_company[df_company["Phân loại"].astype(str).str.strip() == "SI"].copy()
        
        if search_kw:
            mask = (
                df_si["Mã SP"].astype(str).str.contains(search_kw, case=False, na=False)
                | df_si["Tên SP (PO)"].astype(str).str.contains(search_kw, case=False, na=False)
            )
            df_si = df_si[mask]

        if df_si.empty:
            st.info("Không tìm thấy Mã SP nào phù hợp với từ khóa.")
        else:
            list_ma_sp_hien_thi = df_si["Mã SP"].dropna().astype(str).tolist()
            
            cac_ma_sp_chon = st.multiselect(
                "🎯 Hoặc chọn trực tiếp các Mã SP muốn xem (Bỏ trống để xem tất cả):",
                options=list_ma_sp_hien_thi,
                default=[]
            )
            
            if cac_ma_sp_chon:
                df_si = df_si[df_si["Mã SP"].astype(str).isin(cac_ma_sp_chon)]
                
            st.write(f"Đang hiển thị **{len(df_si)}** mã sản phẩm. Bấm vào để xem và nhập liệu:")
            
            for _, row_sp in df_si.iterrows():
                ma_sp = row_sp["Mã SP"]
                ten_sp = row_sp["Tên SP (PO)"] if pd.notna(row_sp["Tên SP (PO)"]) else "Chưa có tên"
                
                df_detail = df_company[df_company["Mã SP"] == ma_sp]
                
                with st.expander(f"📦 {ma_sp} — {ten_sp}", expanded=False):
                    df_po = df_detail[df_detail["Phân loại"].astype(str).str.strip() == "-"].copy()
                    list_po = df_po["Số PO"].dropna().astype(str).unique().tolist()
                    
                    # Bảng 1: Thông tin PO
                    if not df_po.empty:
                        st.markdown("**📌 Thông tin Đơn hàng (PO):**")
                        df_giao_all = df_detail[
                            (df_detail["Phân loại"].astype(str).str.strip() != "SI") & 
                            (df_detail["Phân loại"].astype(str).str.strip() != "-")
                        ].copy()
                        
                        tong_giao_dict = df_giao_all.groupby("Số PO")["Số lượng giao"].sum().to_dict()
                        df_po_display = df_po.copy()
                        df_po_display["Số lượng trên ĐH"] = pd.to_numeric(df_po_display["Số lượng trên ĐH"], errors='coerce').fillna(0)
                        df_po_display["SL Đã Giao"] = df_po_display["Số PO"].map(tong_giao_dict).fillna(0)
                        df_po_display["Còn Lại"] = df_po_display["Số lượng trên ĐH"] - df_po_display["SL Đã Giao"]
                        df_po_display["Ngày giao hàng"] = pd.to_datetime(df_po_display["Ngày giao hàng"]).dt.strftime('%d/%m/%Y')
                        
                        st.dataframe(
                            df_po_display[["Số PO", "Ngày giao hàng", "Số lượng trên ĐH", "SL Đã Giao", "Còn Lại", "Đơn giá", "Hoàn thành"]],
                            use_container_width=True,
                            hide_index=True,
                            column_config={
                                "Số lượng trên ĐH": st.column_config.NumberColumn("SL trên ĐH", format="%d"),
                                "SL Đã Giao": st.column_config.NumberColumn("Đã giao", format="%d"),
                                "Còn Lại": st.column_config.NumberColumn("Còn lại", format="%d"),
                            }
                        )
                    
                    # Bảng 2: Lịch sử giao hàng (Sửa trực tiếp)
                    st.markdown("**🚚 Lịch sử Giao hàng (Sửa trực tiếp trên bảng):**")
                    if df_giao_all.empty:
                        st.info("Chưa có đợt giao hàng nào.")
                    else:
                        for po in list_po:
                            df_giao_po = df_giao_all[df_giao_all["Số PO"].astype(str) == str(po)].copy()
                            if not df_giao_po.empty:
                                st.markdown(f"**PO: {po}**")
                                
                                df_edit = df_giao_po.copy()
                                df_edit["Ngày giao hàng"] = pd.to_datetime(df_edit["Ngày giao hàng"], errors='coerce').dt.strftime('%d/%m/%Y').fillna("")
                                df_edit["Số lượng giao"] = pd.to_numeric(df_edit["Số lượng giao"], errors='coerce').fillna(0)
                                df_edit.insert(0, "Xóa", False)
                                
                                # Sử dụng clean_key để làm sạch key
                                safe_key = clean_key(f"{ma_sp}_{po}")
                                
                                edited_df = st.data_editor(
                                    df_edit[["Xóa", "Ngày giao hàng", "Số lượng giao", "Ghi chú"]],
                                    column_config={
                                        "Xóa": st.column_config.CheckboxColumn("Xóa?", default=False),
                                        "Ngày giao hàng": st.column_config.TextColumn("Ngày giao (DD/MM/YYYY)"),
                                        "Số lượng giao": st.column_config.NumberColumn("SL Giao", format="%d"),
                                        "Ghi chú": st.column_config.TextColumn("Ghi chú"),
                                    },
                                    hide_index=True,
                                    use_container_width=True,
                                    key=f"editor_{safe_key}"
                                )
                                
                                # Nút lưu cho từng PO
                                if st.button(f"💾 Lưu thay đổi cho PO {po}", key=f"btn_update_{safe_key}"):
                                    try:
                                        wb = openpyxl.load_workbook(file_path)
                                        ws = wb[cong_ty]
                                        
                                        header_row = None
                                        for i in range(1, 20):
                                            for j in range(1, ws.max_column + 1):
                                                if ws.cell(row=i, column=j).value and "Số PO" in str(ws.cell(row=i, column=j).value):
                                                    header_row = i
                                                    break
                                            if header_row: break
                                        
                                        for idx, row in edited_df.iloc[::-1].iterrows():
                                            if row["Xóa"]:
                                                target_ngay_old = df_edit.loc[idx, "Ngày giao hàng"]
                                                target_sl_old = df_edit.loc[idx, "Số lượng giao"]
                                                for r in range(ws.max_row, header_row, -1):
                                                    val_po = ws.cell(row=r, column=5).value
                                                    val_ngay = ws.cell(row=r, column=4).value
                                                    val_sl = ws.cell(row=r, column=8).value
                                                    if str(val_po) == str(po) and pd.notna(val_ngay) and pd.notna(val_sl):
                                                        try: date_str = val_ngay.strftime('%d/%m/%Y')
                                                        except: date_str = str(val_ngay)
                                                        if date_str == target_ngay_old and float(val_sl) == float(target_sl_old):
                                                            ws.delete_rows(r, 1)
                                                            break
                                            else:
                                                target_ngay_old = df_edit.loc[idx, "Ngày giao hàng"]
                                                target_sl_old = df_edit.loc[idx, "Số lượng giao"]
                                                new_ngay_str = row["Ngày giao hàng"]
                                                new_sl = row["Số lượng giao"]
                                                new_ghi_chu = row["Ghi chú"]
                                                
                                                if new_ngay_str != target_ngay_old or float(new_sl) != float(target_sl_old) or new_ghi_chu != df_edit.loc[idx, "Ghi chú"]:
                                                    for r in range(ws.max_row, header_row, -1):
                                                        val_po = ws.cell(row=r, column=5).value
                                                        val_ngay = ws.cell(row=r, column=4).value
                                                        val_sl = ws.cell(row=r, column=8).value
                                                        if str(val_po) == str(po) and pd.notna(val_ngay) and pd.notna(val_sl):
                                                            try: date_str = val_ngay.strftime('%d/%m/%Y')
                                                            except: date_str = str(val_ngay)
                                                            if date_str == target_ngay_old and float(val_sl) == float(target_sl_old):
                                                                try:
                                                                    new_date = datetime.strptime(new_ngay_str, '%d/%m/%Y')
                                                                    ws.cell(row=r, column=4).value = new_date
                                                                except:
                                                                    ws.cell(row=r, column=4).value = new_ngay_str
                                                                ws.cell(row=r, column=8).value = new_sl
                                                                ws.cell(row=r, column=11).value = new_ghi_chu if new_ghi_chu != "None" else ""
                                                                break
                                        
                                        wb.save(file_path)
                                        st.success(f"Đã cập nhật thay đổi cho PO {po}!")
                                        st.session_state.pop("all_sheets", None)
                                        st.session_state.pop("visible_sheets", None)
                                        st.rerun()
                                    except Exception as e:
                                        st.error(f"Lỗi khi cập nhật: {e}")
                                st.divider()

                    # Form nhập liệu mới
                    col_tao_po, col_giao_hang = st.columns(2)
                    with col_tao_po:
                        st.markdown("**📝 Tạo PO Mới**")
                        with st.form(key=f"form_po_{clean_key(ma_sp)}"):
                            new_po = st.text_input("Nhập Số PO mới", placeholder="VD: DH26.11.001", key=f"newpo_{clean_key(ma_sp)}")
                            ngay_po = st.date_input("Ngày PO", value=date.today(), key=f"ngaypo_{clean_key(ma_sp)}", format="DD/MM/YYYY")
                            sl_dat = st.number_input("Số lượng khách đặt", min_value=0.0, step=100.0, key=f"sldat_{clean_key(ma_sp)}")
                            don_gia = st.number_input("Đơn giá", min_value=0.0, step=100.0, key=f"gia_{clean_key(ma_sp)}")
                            btn_tao_po = st.form_submit_button("💾 Lưu PO Mới", use_container_width=True)
                            
                            if btn_tao_po:
                                if not new_po:
                                    st.error("Vui lòng nhập Số PO mới.")
                                elif sl_dat <= 0:
                                    st.error("Số lượng khách đặt phải lớn hơn 0.")
                                else:
                                    try:
                                        dm.append_delivery(
                                            file_path=file_path, sheet_name=cong_ty,
                                            ma_sp=ma_sp, so_po=new_po,
                                            ngay_giao=datetime.combine(ngay_po, datetime.min.time()),
                                            so_luong_giao=0, don_gia=don_gia, ghi_chu="",
                                            sl_tren_dh=sl_dat, ten_sp=None
                                        )
                                        st.success(f"✅ Đã tạo PO {new_po} thành công!")
                                        st.session_state.pop("all_sheets", None)
                                        st.session_state.pop("visible_sheets", None)
                                        st.rerun()
                                    except Exception as e:
                                        st.error(f"Lỗi: {e}")

                    with col_giao_hang:
                        st.markdown("**🚚 Nhập Giao Hàng**")
                        if not list_po:
                            st.info("Chưa có PO nào để giao. Vui lòng tạo PO mới ở cột bên trái.")
                        else:
                            with st.form(key=f"form_giao_{clean_key(ma_sp)}"):
                                po_chon = st.selectbox("Chọn PO cần giao", list_po, key=f"chonpo_{clean_key(ma_sp)}")
                                ngay_giao = st.date_input("Ngày giao hàng", value=date.today(), key=f"ngaygiao_{clean_key(ma_sp)}", format="DD/MM/YYYY")
                                so_luong_giao = st.number_input("Số lượng giao thực tế", min_value=0.0, step=1.0, key=f"slgiao_{clean_key(ma_sp)}")
                                ghi_chu = st.text_input("Ghi chú", key=f"ghichu_{clean_key(ma_sp)}")
                                btn_giao_hang = st.form_submit_button("💾 Lưu Giao Hàng", use_container_width=True)
                                
                                if btn_giao_hang:
                                    if so_luong_giao <= 0:
                                        st.error("Số lượng giao phải lớn hơn 0.")
                                    else:
                                        try:
                                            dm.append_delivery(
                                                file_path=file_path, sheet_name=cong_ty,
                                                ma_sp=ma_sp, so_po=po_chon,
                                                ngay_giao=datetime.combine(ngay_giao, datetime.min.time()),
                                                so_luong_giao=so_luong_giao, don_gia=None,
                                                ghi_chu=ghi_chu, sl_tren_dh=None, ten_sp=None
                                            )
                                            st.success(f"✅ Đã ghi nhận giao hàng cho PO {po_chon}!")
                                            st.session_state.pop("all_sheets", None)
                                            st.session_state.pop("visible_sheets", None)
                                            st.rerun()
                                        except Exception as e:
                                            st.error(f"Lỗi: {e}")
    else:
        st.warning(f"Sheet **{cong_ty}** không có cấu trúc đơn hàng hoặc đang trống.")


# ---------------------------------------------------------
# TAB 2: LỊCH SỬ GIAO HÀNG THEO NGÀY
# ---------------------------------------------------------
with tab_loc_ngay:
    st.header("📅 Tra cứu Lịch sử Giao hàng theo Ngày")
    
    col_loc_ngay, col_loc_cong_ty = st.columns(2)
    
    with col_loc_ngay:
        ngay_loc = st.date_input("1. Chọn ngày cần tra cứu", value=date.today(), format="DD/MM/YYYY")
        ngay_loc_str = ngay_loc.strftime('%d/%m/%Y')
        
    with col_loc_cong_ty:
        list_cong_ty_loc = ["Tất cả"] + visible_sheets
        cong_ty_loc = st.selectbox("2. Lọc theo Công ty", list_cong_ty_loc)
    
    frames_giao_hang = []
    for sheet_name, df in all_data.items():
        if cong_ty_loc != "Tất cả" and sheet_name != cong_ty_loc:
            continue
            
        if isinstance(df, pd.DataFrame) and "Ngày giao hàng" in df.columns:
            df_giao = df[
                (df["Phân loại"].astype(str).str.strip() != "SI") & 
                (df["Phân loại"].astype(str).str.strip() != "-")
            ].copy()
            
            if not df_giao.empty:
                df_giao["Ngày_str"] = pd.to_datetime(df_giao["Ngày giao hàng"], errors='coerce').dt.strftime('%d/%m/%Y')
                df_giao_loc = df_giao[df_giao["Ngày_str"] == ngay_loc_str]
                
                if not df_giao_loc.empty:
                    df_giao_loc["_sheet"] = sheet_name
                    frames_giao_hang.append(df_giao_loc)
    
    if not frames_giao_hang:
        st.info(f"Không có đợt giao hàng nào vào ngày {ngay_loc_str} (Công ty: {cong_ty_loc}).")
    else:
        df_ket_qua = pd.concat(frames_giao_hang, ignore_index=True)
        df_ket_qua = df_ket_qua[["_sheet", "Mã SP", "Tên SP (PO)", "Số PO", "Số lượng giao", "Ghi chú"]]
        df_ket_qua.rename(columns={"_sheet": "Công ty"}, inplace=True)
        
        st.success(f"Tìm thấy **{len(df_ket_qua)}** đợt giao hàng vào ngày {ngay_loc_str}.")
        st.dataframe(
            df_ket_qua,
            use_container_width=True,
            hide_index=True,
            column_config={
                "Số lượng giao": st.column_config.NumberColumn("SL Giao", format="%d"),
            }
        )
        
        st.markdown("---")
        st.markdown("**📊 Tổng kết trong ngày:**")
        tong_sp = df_ket_qua["Mã SP"].nunique()
        tong_po = df_ket_qua["Số PO"].nunique()
        tong_sl = df_ket_qua["Số lượng giao"].sum()
        
        c1, c2, c3 = st.columns(3)
        c1.metric("Số Mã SP đã giao", tong_sp)
        c2.metric("Số PO đã giao", tong_po)
        c3.metric("Tổng SL đã giao", f"{tong_sl:,.0f}")


# ---------------------------------------------------------
# TAB 3: TỔNG HỢP KPI
# ---------------------------------------------------------
with tab_tong_hop:
    st.header("📊 Báo cáo Tổng hợp")
    
    frames = []
    for name, df in all_data.items():
        if isinstance(df, pd.DataFrame) and "Số PO" in df.columns:
            sub = df.dropna(subset=["Số PO"]).copy()
            frames.append(sub)
            
    if frames:
        df_all = pd.concat(frames, ignore_index=True)
        df_dat = df_all[df_all["Phân loại"].astype(str).str.strip() == "-"]
        info_dat = df_dat.groupby(["Số PO", "_sheet"]).agg(
            MaSP=("Mã SP", "first"),
            TenSP=("Tên SP (PO)", "first"),
            SLDat=("Số lượng trên ĐH", "first"),
            DonGia=("Đơn giá", "first"),
            TrangThai=("Hoàn thành", "first")
        ).reset_index()
        
        df_giao = df_all[
            (df_all["Phân loại"].astype(str).str.strip() != "SI") & 
            (df_all["Phân loại"].astype(str).str.strip() != "-")
        ]
        info_giao = df_giao.groupby("Số PO").agg(
            SLGiao=("Số lượng giao", "sum")
        ).reset_index()
        
        summary = pd.merge(info_dat, info_giao, on="Số PO", how="left")
        summary["SLGiao"] = summary["SLGiao"].fillna(0)
        summary["ConLai"] = summary["SLDat"] - summary["SLGiao"]
        summary["% Hoàn thành"] = (summary["SLGiao"] / summary["SLDat"].replace(0, pd.NA) * 100).round(1).fillna(0)
        
        c1, c2, c3 = st.columns(3)
        c1.metric("Tổng số PO", len(summary))
        c2.metric("Tổng SL Đặt", f"{summary['SLDat'].sum():,.0f}")
        c3.metric("Tổng SL Giao", f"{summary['SLGiao'].sum():,.0f}")
        
        st.dataframe(
            summary[["_sheet", "MaSP", "Số PO", "SLDat", "SLGiao", "ConLai", "% Hoàn thành", "TrangThai"]],
            use_container_width=True,
            hide_index=True,
            column_config={
                "_sheet": "Công ty",
                "MaSP": "Mã SP",
                "Số PO": "Số PO",
                "SLDat": st.column_config.NumberColumn("SL Đặt", format="%d"),
                "SLGiao": st.column_config.NumberColumn("SL Giao", format="%d"),
                "ConLai": st.column_config.NumberColumn("Còn lại", format="%d"),
                "% Hoàn thành": st.column_config.ProgressColumn("% HT", min_value=0, max_value=100, format="%.1f%%"),
                "TrangThai": "Trạng thái"
            }
        )
    else:
        st.info("Chưa có dữ liệu.")
